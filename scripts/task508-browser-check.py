#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Task 508 — браузерная проверка (Playwright; паттерн task505/506).

Заявка: «Добавь сохранение в архив журнала кто последний вносил
правки в каждый кабель, по столбцу name в таблице users файла
KIP8_Access. И удалённый кабель так же должен сохраняться с
пометкой как архивный по причине "Удалён", кто удалил, и не должен
отображаться в распечатываемом или сохраняемом в PDF файле, а в
excel отображается.»

МОК Apps Script имитирует сервер ПОСЛЕ обновления CableJournal.gs
(Task 508): 6 записей (4 активные + 2 архивные «Удалён», порядок
строк перемешан), getColumns — 20 колонок (5 новых readOnly),
deleteRow — СОСТОЯТЕЛЬНЫЙ: помечает запись архивной (как сервер).

Проверяется:
  A) десктоп тёмная: список БЕЗ архивных (4 карточки); _printModel —
     4 строки (архив не печатается); _excelModel — 6 (архив в Excel);
     диалог «записей: 4», печатный лист/предпросмотр — 4 строки;
     Excel (openpyxl): блок «Архив» K–O, 7 объединений, «Удалён»/
     кем/даты/редактор в данных, 6 строк; PDF — структурно;
     e2e-удаление: submitDelete → тост «в архиве журнала», строка
     исчезает из списка (load), но остаётся в Excel-модели;
  B) светлая: диалог «записей: 3» (после удаления), 3 строки;
  C) мобайл: диалог на весь экран, 3 строки.
"""
import json
import os
import re
import threading
from http.server import HTTPServer, SimpleHTTPRequestHandler
from urllib.parse import unquote

from playwright.sync_api import sync_playwright

REPO = '/home/z/my-project/kip8test'
PORT = 8994
SHOT_DIR = '/home/z/my-project/download/kip8test-task508'
os.makedirs(SHOT_DIR, exist_ok=True)

PASS = 0
FAIL = 0


def check(name, ok, extra=''):
    global PASS, FAIL
    mark = 'OK ' if ok else 'FAIL'
    print('  [%s] %s%s' % (mark, name,
                           (' — %s' % (extra,))[:200] if extra and not ok else ''))
    if ok:
        PASS += 1
    else:
        FAIL += 1


def shot(page, name):
    try:
        page.screenshot(path=os.path.join(SHOT_DIR, name))
    except Exception as e:
        print('  (скриншот %s не сохранён: %s)' % (name, e))


# ---- мок: колонки (15 + 5 Task 508 readOnly) ----
COLUMNS = [
    {'key': 'num', 'label': '№ п/п', 'type': 'number', 'readOnly': True},
    {'key': 'designation', 'label': 'Обозначение кабеля, провода',
     'type': 'text'},
    {'key': 'start', 'label': 'Начало трассы', 'type': 'text'},
    {'key': 'end', 'label': 'Конец трассы', 'type': 'text'},
    {'key': 'section', 'label': 'Участок трассы кабеля, провода',
     'type': 'text'},
    {'key': 'mark_project', 'label': 'Марка (по проекту)', 'type': 'text'},
    {'key': 'cores_project',
     'label': 'Число и сечение жил (по проекту)', 'type': 'text'},
    {'key': 'length_project', 'label': 'Длина, м (по проекту)',
     'type': 'number'},
    {'key': 'mark_actual', 'label': 'Марка (проложен)', 'type': 'text'},
    {'key': 'cores_actual',
     'label': 'Число и сечение жил (проложен)', 'type': 'text'},
    {'key': 'length_actual', 'label': 'Длина, м (проложен)',
     'type': 'number'},
    {'key': 'added_at', 'label': 'Дата добавления', 'type': 'date',
     'readOnly': True},
    {'key': 'department', 'label': 'Отделение, производство',
     'type': 'text'},
    {'key': 'purpose', 'label': 'Назначение кабеля', 'type': 'text'},
    {'key': 'project_no', 'label': '№ проекта', 'type': 'text'},
    # Task 508 — колонки архива (readOnly, сервер пишет)
    {'key': 'last_editor', 'label': 'Последний редактор', 'type': 'text',
     'readOnly': True},
    {'key': 'edited_at', 'label': 'Дата правки', 'type': 'date',
     'readOnly': True},
    {'key': 'archived', 'label': 'Архив', 'type': 'text',
     'readOnly': True},
    {'key': 'archived_by', 'label': 'Кем удалён', 'type': 'text',
     'readOnly': True},
    {'key': 'archived_at', 'label': 'Дата удаления', 'type': 'date',
     'readOnly': True},
]

EDITOR_A = 'Петров Пётр Петрович'
EDITOR_B = 'Тестов Тест Тестович'
EDITOR_C = 'Админов Админ Админович'


def _row(rw, num, des, dep, editor, edited, arch=None):
    r = {'_row': rw, 'num': num, 'designation': des,
         'start': 'ЩК-4, ШУ-%d' % (num % 7 + 1),
         'end': 'КС-%d, отм. %d,0 м.' % (num % 4 + 1, num % 18),
         'section': ['по коробу', 'по лотку', 'фальшпол'][num % 3],
         'mark_project': 'КВБбШЭнг-LS', 'cores_project': '%dх1' % (num + 10),
         'length_project': num * 11, 'mark_actual': 'КВБбШЭнг-LS',
         'cores_actual': '%dх1' % (num + 10), 'length_actual': num * 11,
         'department': dep, 'purpose': 'Контрольный', 'project_no': 'АН-603',
         'added_at': 'Нет данных', 'last_editor': editor,
         'edited_at': edited}
    if arch:
        r['_archived'] = True
        r['archived'] = 'Удалён'
        r['archived_by'] = arch
        r['archived_at'] = '2026-10-%02dT09:30:00.000Z' % (10 + num % 3)
    return r


# 6 записей, ПОРЯДОК ПЕРЕМЕШАН (не № п/п): активные №1/№3/№4/№6,
# архивные №2 («Удалён», Тестов) и №5 («Удалён», Админов)
ROWS = [
    _row(2, 1, '1/7', 'Водоподготовка', EDITOR_A, '2026-10-10T06:00:00'),
    _row(5, 3, '5/2', 'Отделение №2', EDITOR_B, '2026-10-11T06:00:00'),
    _row(3, 2, '2/1', 'Водоподготовка', EDITOR_A, '2026-10-09T06:00:00',
         arch=EDITOR_B),
    _row(6, 4, '4/3', 'МДИ', EDITOR_C, '2026-10-12T06:00:00'),
    _row(7, 5, '5/4', 'МДИ', EDITOR_B, '2026-10-08T06:00:00', arch=EDITOR_C),
    _row(8, 6, '6/5', 'Склад', EDITOR_A, '2026-10-07T06:00:00'),
]


def mock_response(action, body):
    if action == 'getCurrentUser':
        return {'ok': True, 'data': {'userId': 1, 'email': 'user@test.local',
                                     'role': 'Админ'}}
    if action == 'getMyAccess':
        return {'ok': True, 'data': {'role': 'Админ', 'found': True,
                'permissions': {}}}
    if action == 'heartbeat':
        return {'ok': True, 'data': {'ok': True}}
    if action == 'cableJournal.getColumns':
        return {'ok': True, 'data': {'columns': COLUMNS, 'canEdit': True}}
    if action == 'cableJournal.list':
        active = [r for r in ROWS if not r.get('_archived')]
        return {'ok': True, 'data': {'rows': ROWS, 'total': len(active),
                                     'archivedTotal': len(ROWS) - len(active),
                                     'canEdit': True}}
    if action == 'cableJournal.deleteRow':
        # Task 508: как обновлённый сервер — помечаем архивной
        rw = (body or {}).get('row')
        for r in ROWS:
            if r['_row'] == rw:
                r['_archived'] = True
                r['archived'] = 'Удалён'
                r['archived_by'] = EDITOR_B
                r['archived_at'] = '2026-10-12T08:00:00.000Z'
                return {'ok': True, 'data': {'ok': True, 'num': r['num'],
                                             'archived': True}}
        return {'ok': False, 'error': 'Строка не найдена'}
    if action == 'cableJournal.updateRow':
        return {'ok': True, 'data': {'ok': True, 'row': (body or {}).get('row'),
                                     'num': 1, 'updated': []}}
    return {'ok': True, 'data': {'ok': True}}


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def setup_routes(ctx):
    def handle(route, request):
        action = ''
        if 'action=' in request.url:
            action = unquote(request.url.split('action=')[1].split('&')[0])
        body = {}
        if request.post_data:
            try:
                body = json.loads(request.post_data)
            except Exception:
                body = {}
        return route.fulfill(status=200,
            content_type='application/json; charset=utf-8',
            body=json.dumps(mock_response(action, body), ensure_ascii=False))
    ctx.route('**/exec?**', handle)
    ctx.route('**script.google.com/**', handle)
    ctx.route('**raw.githubusercontent.com/**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))
    ctx.route('**calendar.legalic.ru/**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))
    ctx.route('**isdayoff.ru**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))


def seed_ls(theme):
    return ("try{localStorage.clear()}catch(e){};" +
            "localStorage.setItem('kip8test:kip8_session_token','bc-t508');" +
            "localStorage.setItem('kip8test:app-theme','" + theme + "');")


STATE_JS = r"""(() => {
    const btn = document.getElementById('cjPrintBtn');
    const dlg = document.getElementById('cjPrintPrevModal');
    const st = document.getElementById('cjPrintStyle');
    const sheet = document.getElementById('wsPrintSheet');
    const out = {
        btnExists: !!btn,
        dlg: !!dlg,
        styleInjected: !!st,
        cards: document.querySelectorAll('#cjList .cj-item').length,
        groups: document.querySelectorAll('#cjList .cj-group').length,
        sheetExists: !!sheet,
        sheetRows: sheet ? sheet.querySelectorAll(
            '.cjpr-table tbody tr').length : 0
    };
    if (dlg) {
        out.sub = (dlg.querySelector('.wspprev-sub') || {}).textContent || '';
        out.btns = [...dlg.querySelectorAll('.wspprev-btn')].map(
            b => b.textContent.trim());
    }
    return out;})"""

MODELS_JS = r"""(() => {
    // KipCableJournal — глобальная ЛЕКСИЧЕСКАЯ binding (const),
    // НЕ свойство window: обращаемся по имени
    const CJ = KipCableJournal;
    const pm = CJ._printModel();
    const em = CJ._excelModel();
    return {
        printRows: pm.rows.length,
        printDes: pm.rows.map(r => r.cells.designation).join(','),
        excelRows: em.rows.length,
        excelArchived: em.archived,
        excelDes: em.rows.map(r => r.cells.designation).join(','),
        archivedRows: (CJ._archivedRows || []).length
    };})()"""

FRAME_JS = r"""(() => {
    const f = document.querySelector('#cjPrintPrevModal .wspprev-frame');
    if (!f) return {missing: true};
    try {
        const d = f.contentDocument;
        if (!d) return {noDoc: true};
        const des = [];
        d.querySelectorAll('.cjpr-page tbody tr').forEach(tr => {
            des.push((tr.cells[0] || {}).textContent || '');
        });
        return {
            pages: d.querySelectorAll('.cjpr-page').length,
            rows: d.querySelectorAll('.cjpr-page tbody tr').length,
            des: des.join(',')
        };
    } catch (e) { return {err: String(e)}; }})"""


def collect_errors(page, errors):
    page.on('pageerror', lambda e: errors.append('pageerror: ' + str(e)))
    # «Failed to load resource: 429» — сетевой шум картинок Google
    # Drive/lh3 (KipPreload миниатюры; граничат по IP-рейт-лимиту) —
    # НЕ ошибка JavaScript: такие строки не учитываем
    page.on('console', lambda m: errors.append('console: ' + m.text)
            if m.type == 'error' and
            'Failed to load resource' not in m.text else None)


def open_cj(page):
    page.goto('http://localhost:%d/index.html' % PORT)
    page.wait_for_timeout(2500)
    page.evaluate("navigateTo('cable-journal-edit')")
    page.wait_for_timeout(1800)


def save_bytes(download, name):
    path = download.path()
    data = open(path, 'rb').read()
    out = os.path.join(SHOT_DIR, name)
    with open(out, 'wb') as f:
        f.write(data)
    return out, data, download.suggested_filename


def verify_xlsx(path):
    """Excel: блок «Архив» K–O + архивные строки в данных."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb['Кабельный журнал']
    res = []
    merges = sorted(str(m) for m in ws.merged_cells.ranges)
    need = sorted(['A1:A2', 'B1:B2', 'C1:C2', 'D1:D2', 'E1:G1', 'H1:J1',
                   'K1:O1'])
    res.append(('объединения шапки + K1:O1 «Архив» (7 шт)',
                merges == need, merges))
    res.append(('K1 «Архив» — заголовок блока', ws['K1'].value == 'Архив',
                ws['K1'].value))
    res.append(('подколонки блока: K2/L2/M2/N2/O2',
                ws['K2'].value == 'Статус' and
                ws['L2'].value == 'Кем удалён' and
                ws['M2'].value == 'Дата удаления' and
                ws['N2'].value == 'Последний редактор' and
                ws['O2'].value == 'Дата правки',
                (ws['K2'].value, ws['L2'].value, ws['M2'].value,
                 ws['N2'].value, ws['O2'].value)))
    # данные: r3=№1 (актив), r4=№2 (АРХИВ), r5=№3 …
    res.append(('порядок строк: r3=1/7, r4=2/1, r5=5/2',
                ws['A3'].value == '1/7' and ws['A4'].value == '2/1' and
                ws['A5'].value == '5/2',
                (ws['A3'].value, ws['A4'].value, ws['A5'].value)))
    res.append(('архивная строка №2: K4 «Удалён»',
                ws['K4'].value == 'Удалён', ws['K4'].value))
    res.append(('архивная строка №2: L4 кем удалён',
                ws['L4'].value == EDITOR_B, ws['L4'].value))
    res.append(('архивная строка №2: M4 дата «ДД.ММ.ГГГГ»',
                ws['M4'].value == '12.10.2026', ws['M4'].value))
    res.append(('редактор активной строки: N3', ws['N3'].value == EDITOR_A,
                ws['N3'].value))
    res.append(('дата правки активной строки: O3',
                ws['O3'].value == '10.10.2026', ws['O3'].value))
    res.append(('активная строка БЕЗ статуса: K3 пусто',
                ws['K3'].value in (None, ''), ws['K3'].value))
    res.append(('архивная строка №5: K7 «Удалён», L7 кем',
                ws['K7'].value == 'Удалён' and ws['L7'].value == EDITOR_C,
                (ws['K7'].value, ws['L7'].value)))
    n_arch = sum(1 for r in range(3, ws.max_row + 1)
                 if ws.cell(row=r, column=11).value == 'Удалён')
    res.append(('архивных строк в Excel: 2 (в печати их нет!)',
                n_arch == 2, n_arch))
    n_rows = sum(1 for r in range(3, ws.max_row + 1)
                 if any(c.value not in (None, '')
                        for c in ws[r][:10]))
    res.append(('всего строк данных: 6 (Excel показывает архив)',
                n_rows == 6, n_rows))
    res.append(('Print_Titles 1:2 жив',
                ws.print_title_rows in ('1:2', '$1:$2'),
                ws.print_title_rows))
    res.append(('закрепление A3 живо', ws.freeze_panes == 'A3',
                ws.freeze_panes))
    return res


def main():
    os.chdir(REPO)
    server = HTTPServer(('127.0.0.1', PORT), QuietHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    with sync_playwright() as p:
        browser = p.chromium.launch()

        # ======== A. ДЕСКТОП, тёмная ========
        print('A) ДЕСКТОП 1280x720, тёмная — архив: список/модели/Excel/удаление')
        ctx = browser.new_context(viewport={'width': 1280, 'height': 720},
                                  accept_downloads=True)
        ctx.add_init_script(seed_ls('dark'))
        setup_routes(ctx)
        page = ctx.new_page()
        errors = []
        collect_errors(page, errors)
        open_cj(page)

        st = page.evaluate(STATE_JS)
        check('список БЕЗ архивных: 4 карточки (архив скрыт)',
              st['cards'] == 4, st)
        check('группы: 4 отделения', st['groups'] == 4, st['groups'])
        shot(page, 'a1-list-no-archive.png')

        md = page.evaluate(MODELS_JS)
        check('_printModel: 4 строки (архив НЕ печатается)',
              md['printRows'] == 4, md)
        check('_printModel: обозначения без архивных',
              md['printDes'] == '1/7,5/2,4/3,6/5', md['printDes'])
        check('_excelModel: 6 строк (архив в Excel)',
              md['excelRows'] == 6 and md['excelArchived'] == 2, md)
        check('_excelModel: сортировка № п/п (архив меж активных)',
              md['excelDes'] == '1/7,2/1,5/2,4/3,5/4,6/5', md['excelDes'])
        check('_archivedRows: 2 записи в состоянии модуля',
              md['archivedRows'] == 2, md['archivedRows'])

        # диалог печати
        page.click('#cjPrintBtn')
        page.wait_for_timeout(1000)
        st = page.evaluate(STATE_JS)
        check('диалог открыт; подзаголовок «записей: 4» (не 6)',
              st['dlg'] and 'записей: 4' in st.get('sub', ''), st.get('sub'))
        check('печатный лист: 4 строки (архив не в печати)',
              st['sheetRows'] == 4, st['sheetRows'])
        fr = page.evaluate(FRAME_JS)
        check('предпросмотр: 4 строки, без архивных обозначений',
              fr.get('rows') == 4 and
              fr.get('des') == '1/7,5/2,4/3,6/5', fr)
        shot(page, 'a2-print-no-archive.png')

        # PDF: структурно (архив не попадает — модель та же)
        try:
            with page.expect_download(timeout=30000) as dl:
                page.click('#cjPrintPrevModal .wspprev-pdf')
            out, data, fname = save_bytes(dl.value, 'task508.pdf')
            check('PDF: имя + %PDF-магия',
                  fname.startswith('Кабельный_журнал_') and
                  fname.endswith('.pdf') and data[:4] == b'%PDF' and
                  len(data) > 2000, (fname, len(data)))
        except Exception as e:
            check('PDF скачан', False, str(e)[:200])
        page.wait_for_timeout(500)

        # Excel: блок «Архив» + архивные строки (openpyxl)
        try:
            with page.expect_download(timeout=30000) as dl:
                page.click('#cjPrintPrevModal .wspprev-xlsx')
            out, data, fname = save_bytes(dl.value, 'task508.xlsx')
            check('Excel: имя + PK-магия',
                  fname.startswith('Кабельный_журнал_') and
                  fname.endswith('.xlsx') and data[:2] == b'PK', fname)
            for name, ok, extra in verify_xlsx(out):
                check('xlsx: ' + name, ok, extra)
        except Exception as e:
            check('Excel скачан', False, str(e)[:300])
        page.wait_for_timeout(500)
        page.keyboard.press('Escape')
        page.wait_for_timeout(400)

        # e2e-удаление: submitDelete → тост «в архиве», строка уходит
        # из списка, но остаётся в Excel-модели (мок пометил архивом)
        page.evaluate("() => { const CJ = KipCableJournal;" +
                      " CJ._viewing = CJ._rows.find(r => r.num === 3);" +
                      " CJ.submitDelete(); }")
        page.wait_for_timeout(1500)
        toast = page.evaluate(
            "() => (document.getElementById('toastMessage') || " +
            "{}).textContent || ''")
        check('тост удаления упоминает архив журнала',
              'удалена' in toast and 'архиве' in toast, toast)
        st = page.evaluate(STATE_JS)
        check('после удаления: карточек 3 (5/2 ушла из списка)',
              st['cards'] == 3, st['cards'])
        md = page.evaluate(MODELS_JS)
        check('после удаления: печать 3 строки',
              md['printRows'] == 3 and
              md['printDes'] == '1/7,4/3,6/5', md)
        check('после удаления: Excel 6 строк (архивная добавлена)',
              md['excelRows'] == 6 and md['excelArchived'] == 3, md)
        shot(page, 'a3-after-delete.png')

        check('0 JS-ошибок (тёмная, все шаги)', len(errors) == 0,
              errors[:4])
        ctx.close()

        # ======== B. ДЕСКТОП, светлая ========
        print('B) ДЕСКТОП, светлая — диалог после удаления')
        ctx2 = browser.new_context(viewport={'width': 1280, 'height': 720},
                                   accept_downloads=True)
        ctx2.add_init_script(seed_ls('light'))
        setup_routes(ctx2)
        page2 = ctx2.new_page()
        errors2 = []
        collect_errors(page2, errors2)
        open_cj(page2)
        st2 = page2.evaluate(STATE_JS)
        check('светлая: 3 карточки (мок с архивной после удаления)',
              st2['cards'] == 3, st2['cards'])
        page2.click('#cjPrintBtn')
        page2.wait_for_timeout(1000)
        st2 = page2.evaluate(STATE_JS)
        check('светлая: диалог «записей: 3» (архив не печатается)',
              st2['dlg'] and 'записей: 3' in st2.get('sub', ''),
              st2.get('sub'))
        fr2 = page2.evaluate(FRAME_JS)
        check('светлая: предпросмотр 3 строки',
              fr2.get('rows') == 3, fr2)
        shot(page2, 'b1-light-dialog.png')
        page2.keyboard.press('Escape')
        page2.wait_for_timeout(300)
        check('0 JS-ошибок (светлая)', len(errors2) == 0, errors2[:4])
        ctx2.close()

        # ======== C. МОБАЙЛ 375x812, тёмная ========
        print('C) МОБАЙЛ 375x812, тёмная — диалог')
        ctx3 = browser.new_context(viewport={'width': 375, 'height': 812},
                                   device_scale_factor=2)
        ctx3.add_init_script(seed_ls('dark'))
        setup_routes(ctx3)
        page3 = ctx3.new_page()
        errors3 = []
        collect_errors(page3, errors3)
        open_cj(page3)
        st3 = page3.evaluate(STATE_JS)
        check('мобайл: список 3 карточки', st3['cards'] == 3, st3['cards'])
        page3.click('#cjPrintBtn')
        page3.wait_for_timeout(1000)
        mst = page3.evaluate(
            r"""(() => {
                const dlg = document.querySelector(
                    '#cjPrintPrevModal .wspprev-dialog');
                if (!dlg) return {dlg: false};
                const r = dlg.getBoundingClientRect();
                return {dlg: true, w: r.width,
                        vw: window.innerWidth};})()""")
        check('мобайл: диалог на весь экран',
              mst.get('dlg') and abs(mst['w'] - 375) < 2, mst)
        fr3 = page3.evaluate(FRAME_JS)
        check('мобайл: предпросмотр 3 строки (архив не печатается)',
              fr3.get('rows') == 3, fr3)
        shot(page3, 'c1-mobile-dialog.png')
        page3.keyboard.press('Escape')
        page3.wait_for_timeout(300)
        check('0 JS-ошибок (мобайл)', len(errors3) == 0, errors3[:4])
        ctx3.close()

        browser.close()

    server.shutdown()
    print()
    print('ИТОГ: %d passed, %d failed' % (PASS, FAIL))
    return 0 if FAIL == 0 else 1


if __name__ == '__main__':
    import sys
    sys.exit(main())

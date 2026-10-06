#!/usr/bin/env python3
"""Task 480: проверка структуры новой Google-таблицы «Перечень КИП ИОС рабочий.xlsx»."""
import openpyxl

PATH = '/home/z/my-project/tmp/new-perench.xlsx'

with open(PATH, 'rb') as f:
    head = f.read(4)
print(f'Первые байты: {head!r}  (xlsx = b"PK\\x03\\x04") -> {"OK" if head == b"PK\\x03\\x04" else "FAIL"}')

wb = openpyxl.load_workbook(PATH, read_only=True, data_only=True)
names = wb.sheetnames
print(f'Листы ({len(names)}): {names}')

REQUIRED = ['Приборы_app', 'Блокировки_app', 'Клапана_app', 'Регуляторы_app']
for req in REQUIRED:
    print(f'  Лист "{req}": {"OK" if req in names else "ОТСУТСТВУЕТ!"}')

# Сверка заголовков листа Приборы_app со структурой, которую ждёт sync-devices.py
if 'Приборы_app' in names:
    ws = wb['Приборы_app']
    headers = [str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows(max_row=1))]
    print(f'Заголовки Приборы_app ({len(headers)}): {headers}')
    # Ключевые колонки, которые использует фронтенд/ППР-индикация (Task 478/479)
    for key in ['ID', 'Наименование', 'Дата', 'В гр. ППР', 'Вид ремонта', 'Период ремонта', 'Изображение']:
        print(f'  Колонка "{key}": {"OK" if key in headers else "ОТСУТСТВУЕТ!"}')
    # Подсчёт строк данных
    cnt = sum(1 for row in ws.iter_rows(min_row=2, max_col=1) if row[0].value is not None)
    print(f'Строк с ID (грубо): {cnt}')
wb.close()

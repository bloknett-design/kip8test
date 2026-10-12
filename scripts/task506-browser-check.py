#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Task 506: browser-check — заявка:
#   «Оформление сделай по шаблону …, без текста "Кабельный журнал
#    отделений: 11 · записей: 159" просто таблицу, шапка таблицы
#    должна распечатываться на каждой странице, просмотр с
#    разбиением листов по страницам, сортировку как в исходном
#    файле.»
# Проверяется #page-cable-journal-edit (мок Apps Script, 30 строк,
# ОТДАННЫХ НЕ ПО ПОРЯДКУ № — модель обязана пересортировать):
#   A) ДЕСКТОП 1280x720, тёмная: кнопка → диалог; iframe:
#      страницы A4 297×210мм (несколько листов!), НА КАЖДОЙ странице
#      своя шапка (thead ×N), шапка-заголовок ОТСУТСТВУЕТ, строк
#      отделений/зебры НЕТ, серые ячейки «По проекту», порядок строк
#      = № п/п (сортировка исходного файла), подзаголовок диалога
#      «записей: 30» без «отделений»; выгрузки PDF (страниц ≥2,
#      /Count) и Excel (openpyxl: объединения, Print_Titles,
#      закрепление, ширины, заливка, порядок); печатный лист в
#      основном документе (window.print); фильтр «из N»; Esc; 0 JS;
#   B) ДЕСКТОП светлая; C) МОБАЙЛ 375x812; D) скриншоты.
import json
import os
import re
import threading
from http.server import HTTPServer, SimpleHTTPRequestHandler
from urllib.parse import unquote

from playwright.sync_api import sync_playwright

PORT = 8984
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOT_DIR = os.path.join(os.path.dirname(REPO), 'download',
                        'kip8test-task506')
os.makedirs(SHOT_DIR, exist_ok=True)

PASS = 0
FAIL = 0


def check(name, cond, extra=''):
    global PASS, FAIL
    ok = bool(cond)
    if ok:
        PASS += 1
    else:
        FAIL += 1
    print(('  + ' if ok else '  X ') + name +
          (('  [' + str(extra)[:300] + ']') if (extra and not ok) else ''))


def shot(page, name):
    try:
        page.screenshot(path=os.path.join(SHOT_DIR, name), full_page=False)
    except Exception as e:
        print('  (скриншот %s не сохранён: %s)' % (name, e))


# ---- мок: 30 строк, порядок ПЕРЕМЕШАН (не № п/п) ----
COLUMNS = [
    {'key': 'designation', 'label': 'Обозначение кабеля, провода',
     'type': 'text'},
    {'key': 'start', 'label': 'Начало трассы', 'type': 'text'},
    {'key': 'end', 'label': 'Конец трассы', 'type': 'text'},
    {'key': 'section', 'label': 'Участок трассы кабеля, провода',
     'type': 'text'},
    {'key': 'mark_project', 'label': 'Марка (по проекту)', 'type': 'text'},
    {'key': 'cores_project',
     'label': 'Число и сечение жил (по проекту)', 'type': 'text'},
    {'key': 'length_project', 'label': 'Длина, м (по проекту)',
     'type': 'number'},
    {'key': 'mark_actual', 'label': 'Марка (проложен)', 'type': 'text'},
    {'key': 'cores_actual',
     'label': 'Число и сечение жил (проложен)', 'type': 'text'},
    {'key': 'length_actual', 'label': 'Длина, м (проложен)',
     'type': 'number'},
    {'key': 'added_at', 'label': 'Дата добавления', 'type': 'text'},
    {'key': 'department', 'label': 'Отделение, производство',
     'type': 'text'},
    {'key': 'purpose', 'label': 'Назначение кабеля', 'type': 'text'},
    {'key': 'project_no', 'label': '№ проекта', 'type': 'text'},
]

ROWS = [
    {'_row': 2, 'num': 1, 'designation': '1/7', 'start': 'ЩК-4, ШУ-1',
     'end': 'ЭЩ №4', 'section': 'по коробу',
     'mark_project': 'КВБбШЭнг-LS', 'cores_project': '14х1',
     'length_project': 60, 'mark_actual': 'КВБбШЭнг-LS',
     'cores_actual': '19х1', 'length_actual': 90,
     'department': 'Водоподготовка', 'purpose': 'Контрольный',
     'project_no': 'АН-603', 'added_at': 'Нет данных'},
    {'_row': 5, 'num': 3, 'designation': '5/2', 'start': 'ЩСУ-2',
     'end': 'Н-1', 'section': 'по лотку',
     'mark_project': 'КГ', 'cores_project': '1.0',
     'length_project': 12, 'mark_actual': 'КГ', 'cores_actual': '3х2,5',
     'length_actual': 12, 'department': 'Отделение №2',
     'purpose': 'Силовой', 'project_no': 'АН-603',
     'added_at': 'Нет данных'},
    {'_row': 3, 'num': 2, 'designation': '2/1', 'start': 'ЩК-4, ШУ-1',
     'end': 'КС-1, отм. 0,0 м., оси 1-2/А-Б', 'section': 'по коробу',
     'mark_project': 'КВБбШЭнг-LS', 'cores_project': '27х1',
     'length_project': 43, 'mark_actual': 'КВБбШЭнг-LS',
     'cores_actual': '27х1', 'length_actual': 43.5,
     'department': 'Водоподготовка', 'purpose': 'Контрольный',
     'project_no': 'АН-603', 'added_at': 'Нет данных'},
]
for i in range(4, 31):
    dep = ['Водоподготовка', 'МДИ', 'КИП и А', 'Склад'][i % 4]
    ROWS.append({
        '_row': i + 1, 'num': i,
        'designation': '%d/%d' % (i, i % 9 + 1),
        'start': 'ЩК-4, ШУ-%d' % (i % 7 + 1),
        'end': 'КС-%d, отм. %d,0 м., оси %d-%d/А-Б' %
               (i % 4 + 1, i % 18, i % 10 + 1, i % 10 + 2),
        'section': ['по коробу', 'по лотку', 'фальшпол', 'труба'][i % 4],
        'mark_project': 'КВБбШЭнг-LS', 'cores_project': '%dх1' % (i % 20 + 2),
        'length_project': (i * 7) % 180 + 5,
        'mark_actual': 'КВБбШЭнг-LS', 'cores_actual': '%dх1' % (i % 20 + 2),
        'length_actual': (i * 7) % 180 + 5,
        'department': dep, 'purpose': 'Контрольный',
        'project_no': 'АН-603', 'added_at': 'Нет данных'})


def mock_response(action, body):
    if action == 'getCurrentUser':
        return {'ok': True, 'data': {'userId': 1, 'email': 'user@test.local',
                                     'role': 'Админ'}}
    if action == 'getMyAccess':
        return {'ok': True, 'data': {'role': 'Админ', 'found': True,
                'permissions': {}}}
    if action == 'heartbeat':
        return {'ok': True, 'data': {'ok': True}}
    if action == 'cableJournal.getColumns':
        return {'ok': True, 'data': {'columns': COLUMNS, 'canEdit': True}}
    if action == 'cableJournal.list':
        return {'ok': True, 'data': {'rows': ROWS, 'total': len(ROWS),
                                     'canEdit': True}}
    return {'ok': True, 'data': {'ok': True}}


class QuietHandler(SimpleHTTPRequestHandler):
    def log_message(self, *args):
        pass


def setup_routes(ctx):
    def handle(route, request):
        action = ''
        if 'action=' in request.url:
            action = unquote(request.url.split('action=')[1].split('&')[0])
        body = {}
        if request.post_data:
            try:
                body = json.loads(request.post_data)
            except Exception:
                body = {}
        return route.fulfill(status=200,
            content_type='application/json; charset=utf-8',
            body=json.dumps(mock_response(action, body), ensure_ascii=False))
    ctx.route('**/exec?**', handle)
    ctx.route('**script.google.com/**', handle)
    ctx.route('**raw.githubusercontent.com/**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))
    ctx.route('**calendar.legalic.ru/**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))
    ctx.route('**isdayoff.ru**', lambda r: r.fulfill(
        status=404, content_type='text/plain', body='nf'))


def seed_ls(theme):
    return ("try{localStorage.clear()}catch(e){};" +
            "localStorage.setItem('kip8test:kip8_session_token','bc-t506');" +
            "localStorage.setItem('kip8test:app-theme','" + theme + "');")


STATE_JS = r"""(() => {
    const btn = document.getElementById('cjPrintBtn');
    const dlg = document.getElementById('cjPrintPrevModal');
    const st = document.getElementById('cjPrintStyle');
    const sheet = document.getElementById('wsPrintSheet');
    const out = {
        btnExists: !!btn,
        dlg: !!dlg,
        styleInjected: !!st,
        sheetExists: !!sheet,
        sheetCls: sheet ? sheet.className : '',
        sheetHasTitle: !!(sheet && sheet.querySelector('.cjpr-title')),
        sheetThRows: sheet ? sheet.querySelectorAll(
            '.cjpr-table thead tr').length : 0,
        sheetRows: sheet ? sheet.querySelectorAll(
            '.cjpr-table tbody tr').length : 0
    };
    if (dlg) {
        const f = dlg.querySelector('.wspprev-frame');
        out.sub = (dlg.querySelector('.wspprev-sub') || {}).textContent || '';
        out.hint = (dlg.querySelector('.wspprev-hint') || {}).textContent || '';
        out.btns = [...dlg.querySelectorAll('.wspprev-btn')].map(
            b => b.textContent.trim());
    }
    return out;})"""

# предпросмотр: страницы, шапки, порядок, серые ячейки, шрифт
FRAME_JS = r"""(() => {
    const f = document.querySelector('#cjPrintPrevModal .wspprev-frame');
    if (!f) return {missing: true};
    try {
        const d = f.contentDocument;
        if (!d) return {noDoc: true};
        const pages = [...d.querySelectorAll('.cjpr-page')];
        const measure = d.getElementById('cjMeasure');
        const font = d.getElementById('wsPrintSheet');
        const first = pages[0];
        const pw = first ? first.getBoundingClientRect().width : 0;
        const ph = first ? first.getBoundingClientRect().height : 0;
        // порядок строк: все первые ячейки страниц подряд
        const des = [];
        d.querySelectorAll('.cjpr-page tbody tr').forEach(tr => {
            des.push((tr.cells[0] || {}).textContent || '');
        });
        return {
            pages: pages.length,
            measureGone: !measure,
            titleGone: !d.querySelector('.cjpr-title, .cjpr-head'),
            grpRows: d.querySelectorAll('tr.cjpr-grp').length,
            zebra: d.querySelectorAll('tr.cjpr-odd, tr.cjpr-even').length,
            theadPerPage: pages.map(p =>
                p.querySelectorAll('thead tr').length),
            thRowspan: first ?
                first.querySelectorAll('th[rowspan="2"]').length : 0,
            thColspan: first ?
                first.querySelectorAll('th[colspan="3"]').length : 0,
            projGray: first ?
                [...first.querySelectorAll('td.cjpr-c-proj')].length : 0,
            grayBg: first ?
                (getComputedStyle(first.querySelector('td.cjpr-c-proj'))
                 .backgroundColor || '') : '',
            fontFamily: font ?
                getComputedStyle(font).fontFamily || '' : '',
            landscape: (d.documentElement.innerHTML
                        .indexOf('A4 landscape') !== -1),
            pageW: Math.round(pw), pageH: Math.round(ph),
            des: des,
            desJoined: des.join(',')
        };
    } catch (e) { return {err: String(e)}; }})"""


def collect_errors(page, errors):
    page.on('pageerror', lambda e: errors.append('pageerror: ' + str(e)))
    page.on('console', lambda m: errors.append('console: ' + m.text)
            if m.type == 'error' else None)


def open_cj(page):
    page.goto('http://localhost:%d/index.html' % PORT)
    page.wait_for_timeout(2500)
    page.evaluate("navigateTo('cable-journal-edit')")
    page.wait_for_timeout(1800)


def save_bytes(download, name):
    path = download.path()
    data = open(path, 'rb').read()
    out = os.path.join(SHOT_DIR, name)
    with open(out, 'wb') as f:
        f.write(data)
    return out, data, download.suggested_filename


def verify_xlsx(path):
    """Глубокая проверка выгруженной книги (openpyxl)."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb['Кабельный журнал']
    res = []
    res.append(('лист «Кабельный журнал»', True, ''))
    merges = sorted(str(m) for m in ws.merged_cells.ranges)
    need = sorted(['A1:A2', 'B1:B2', 'C1:C2', 'D1:D2', 'E1:G1', 'H1:J1'])
    res.append(('объединения шапки A1:A2…D1:D2, E1:G1, H1:J1',
                merges == need, merges))
    res.append(('шапка A1 «Обозначение кабеля, провода»',
                ws['A1'].value == 'Обозначение кабеля, провода',
                ws['A1'].value))
    res.append(('пара E1 «По проекту»', ws['E1'].value == 'По проекту',
                ws['E1'].value))
    res.append(('пара H1 «Проложен»', ws['H1'].value == 'Проложен',
                ws['H1'].value))
    res.append(('подпись F2 «Число и сечение жил»',
                ws['F2'].value == 'Число и сечение жил', ws['F2'].value))
    res.append(('A1 жирная', bool(ws['A1'].font.b), ''))
    res.append(('шрифт Arial Narrow 12',
                ws['A3'].font.name == 'Arial Narrow' and
                ws['A3'].font.sz == 12, (ws['A3'].font.name, ws['A3'].font.sz)))
    res.append(('сетка A3 (thin black)',
                (ws['A3'].border.left.style == 'thin') and
                (ws['A3'].border.left.color.rgb in ('FF000000', '00000000')),
                ws['A3'].border.left.style))
    res.append(('E3 серая F2F2F2',
                (ws['E3'].fill.fgColor.rgb or '').endswith('F2F2F2'),
                ws['E3'].fill.fgColor.rgb))
    res.append(('H3 БЕЗ заливки', ws['H3'].fill.patternType is None,
                ws['H3'].fill.patternType))
    res.append(('закрепление A3', ws.freeze_panes == 'A3',
                ws.freeze_panes))
    res.append(('печать A4 альбомная',
                ws.page_setup.orientation == 'landscape', ''))
    res.append(('Print_Titles 1:2 (шапка на каждой странице)',
                ws.print_title_rows in ('1:2', '$1:$2'),
                ws.print_title_rows))
    res.append(('порядок данных: r3=1/7, r4=2/1, r5=5/2',
                ws['A3'].value == '1/7' and ws['A4'].value == '2/1' and
                ws['A5'].value == '5/2',
                (ws['A3'].value, ws['A4'].value, ws['A5'].value)))
    res.append(('жилы «1.0» → число 1 (шаблон)',
                ws['F5'].value == 1, repr(ws['F5'].value)))
    res.append(('длина G3 — число 60', ws['G3'].value == 60,
                repr(ws['G3'].value)))
    res.append(('нет строк-заголовка (r1 — сразу шапка)',
                ws['A1'].value != 'Кабельный журнал', ws['A1'].value))
    res.append(('ширины колонок шаблона (17/19.71/…)',
                abs((ws.column_dimensions['A'].width or 0) - 17) < 0.2 and
                abs((ws.column_dimensions['G'].width or 0) - 11.43) < 0.2,
                (ws.column_dimensions['A'].width,
                 ws.column_dimensions['G'].width)))
    last = 0
    for r in range(3, ws.max_row + 1):
        if any(c.value not in (None, '') for c in ws[r][:10]):
            last = r
    res.append(('30 строк данных', last == 32, last))
    return res


def main():
    os.chdir(REPO)
    server = HTTPServer(('127.0.0.1', PORT), QuietHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()

    with sync_playwright() as p:
        browser = p.chromium.launch()

        # ======== A. ДЕСКТОП, тёмная ========
        print('A) ДЕСКТОП 1280x720, тёмная — страницы предпросмотра')
        ctx = browser.new_context(viewport={'width': 1280, 'height': 720},
                                  accept_downloads=True)
        ctx.add_init_script(seed_ls('dark'))
        setup_routes(ctx)
        page = ctx.new_page()
        errors = []
        collect_errors(page, errors)
        open_cj(page)

        page.click('#cjPrintBtn')
        page.wait_for_timeout(1200)
        st = page.evaluate(STATE_JS)
        check('диалог открыт, инжект-стиль в <head>',
              st['dlg'] and st['styleInjected'], st)
        check('кнопки: Печать/PDF/Excel/Отмена',
              st.get('btns') == ['Печать', 'Сохранить PDF',
                                 'Сохранить Excel', 'Отмена'],
              st.get('btns'))
        check('подзаголовок «Кабельный журнал — записей: 30» (БЕЗ отделений)',
              'записей: 30' in st.get('sub', '') and
              'отделений' not in st.get('sub', ''), st.get('sub'))
        check('печатный лист в осн. документе: cj-sheet, БЕЗ заголовка',
              st['sheetExists'] and st['sheetCls'] == 'cj-sheet' and
              not st['sheetHasTitle'] and st['sheetThRows'] == 2 and
              st['sheetRows'] == 30, st)

        fr = page.evaluate(FRAME_JS)
        check('предпросмотр: НЕСКОЛЬКО страниц A4',
              fr.get('pages', 0) >= 2, fr.get('pages'))
        check('страница ≈ 297×210мм @96dpi (1123×794px)',
              abs(fr.get('pageW', 0) - 1123) < 26 and
              abs(fr.get('pageH', 0) - 794) < 26,
              (fr.get('pageW'), fr.get('pageH')))
        check('мерная таблица убрана', fr.get('measureGone'))
        check('заголовка-шапки НЕТ («просто таблицу»)',
              fr.get('titleGone'))
        check('шапка НА КАЖДОЙ странице (thead 2×N)',
              fr.get('pages', 0) >= 2 and
              all(t == 2 for t in fr.get('theadPerPage', [])),
              fr.get('theadPerPage'))
        check('шапка шаблона: 4 rowspan + 2 colspan-пары',
              fr.get('thRowspan') == 4 and fr.get('thColspan') == 2,
              (fr.get('thRowspan'), fr.get('thColspan')))
        check('строк отделений и зебры НЕТ',
              fr.get('grpRows') == 0 and fr.get('zebra') == 0,
              (fr.get('grpRows'), fr.get('zebra')))
        check('серые ячейки «По проекту» (rgb(242,242,242))',
              fr.get('projGray', 0) >= 3 and
              '242, 242, 242' in fr.get('grayBg', ''),
              (fr.get('projGray'), fr.get('grayBg')))
        check('шрифт Arial Narrow', 'Arial Narrow' in
              fr.get('fontFamily', ''), fr.get('fontFamily'))
        # сортировка исходного файла: 1/7 → 2/1 → 5/2 (мок перемешан)
        des = fr.get('desJoined', '')
        check('сортировка как в исходном файле (№ п/п: 1/7, 2/1, 5/2…)',
              des.startswith('1/7,2/1,5/2'), des[:60])
        check('альбомная @page в iframe', fr.get('landscape'))
        shot(page, 'a1-preview-pages.png')

        # выгрузка PDF: страниц ≥ 2 (/Count)
        try:
            with page.expect_download(timeout=30000) as dl:
                page.click('#cjPrintPrevModal .wspprev-pdf')
            out, data, fname = save_bytes(dl.value, 'task506.pdf')
            m = re.search(rb'/Type\s*/Pages[^>]*?/Count\s+(\d+)', data)
            npages = int(m.group(1)) if m else 0
            check('PDF: имя + %PDF + ≥2 страниц (шапка на каждой)',
                  fname.startswith('Кабельный_журнал_') and
                  data[:4] == b'%PDF' and len(data) > 5000 and
                  npages >= 2, (fname, npages, len(data)))
        except Exception as e:
            check('PDF скачан', False, str(e)[:200])
        page.wait_for_timeout(600)

        # выгрузка Excel: глубокая проверка openpyxl
        try:
            with page.expect_download(timeout=30000) as dl:
                page.click('#cjPrintPrevModal .wspprev-xlsx')
            out, data, fname = save_bytes(dl.value, 'task506.xlsx')
            check('Excel: имя + PK-магия',
                  fname.startswith('Кабельный_журнал_') and
                  fname.endswith('.xlsx') and data[:2] == b'PK', fname)
            for name, ok, extra in verify_xlsx(out):
                check('xlsx: ' + name, ok, extra)
        except Exception as e:
            check('Excel скачан', False, str(e)[:300])
        page.wait_for_timeout(600)

        # фильтр: «из N»
        page.keyboard.press('Escape')
        page.wait_for_timeout(400)
        page.evaluate("() => { const i = document.getElementById('cjSearchInput');"
                      " i.value = 'Водоподготовка';"
                      " KipCableJournal._applyClientSearch(); }")
        page.wait_for_timeout(400)
        page.click('#cjPrintBtn')
        page.wait_for_timeout(1200)
        st = page.evaluate(STATE_JS)
        fr = page.evaluate(FRAME_JS)
        check('фильтр: подзаголовок «записей: … из 30»',
              'из 30' in st.get('sub', ''), st.get('sub'))
        check('фильтр: в предпросмотре только отфильтрованные',
              0 < fr.get('pages', 0) and
              len(fr.get('des', [])) < 30 and
              all(d for d in fr.get('des', [])),
              len(fr.get('des', [])))
        shot(page, 'a2-filtered.png')

        page.keyboard.press('Escape')
        page.wait_for_timeout(400)
        st = page.evaluate(STATE_JS)
        check('Esc закрыл диалог и снял инжект-стиль',
              not st['dlg'] and not st['styleInjected'], st)
        check('0 JS-ошибок (тёмная, все шаги)', len(errors) == 0,
              errors[:4])
        ctx.close()

        # ======== B. ДЕСКТОП, светлая ========
        print('B) ДЕСКТОП 1280x720, светлая')
        ctx2 = browser.new_context(viewport={'width': 1280, 'height': 720},
                                   accept_downloads=True)
        ctx2.add_init_script(seed_ls('light'))
        setup_routes(ctx2)
        page2 = ctx2.new_page()
        errors2 = []
        collect_errors(page2, errors2)
        open_cj(page2)
        page2.click('#cjPrintBtn')
        page2.wait_for_timeout(1200)
        st = page2.evaluate(STATE_JS)
        fr = page2.evaluate(FRAME_JS)
        check('светлая: диалог + страницы + шапки на каждой',
              st['dlg'] and fr.get('pages', 0) >= 2 and
              all(t == 2 for t in fr.get('theadPerPage', [])),
              fr.get('pages'))
        shot(page2, 'b1-light-preview.png')
        page2.keyboard.press('Escape')
        check('светлая: 0 JS-ошибок', len(errors2) == 0, errors2[:4])
        ctx2.close()

        # ======== C. МОБАЙЛ, тёмная ========
        print('C) МОБАЙЛ 375x812, тёмная')
        ctx3 = browser.new_context(viewport={'width': 375, 'height': 812},
                                   accept_downloads=True)
        ctx3.add_init_script(seed_ls('dark'))
        setup_routes(ctx3)
        page3 = ctx3.new_page()
        errors3 = []
        collect_errors(page3, errors3)
        open_cj(page3)
        btn_visible = page3.evaluate(
            "() => { const b = document.getElementById('cjPrintBtn');"
            " return !!b && b.offsetParent !== null; }")
        check('мобайл: кнопка печати в шапке видна', btn_visible)
        page3.click('#cjPrintBtn')
        page3.wait_for_timeout(1200)
        fr = page3.evaluate(FRAME_JS)
        check('мобайл: диалог открыт, страницы живы',
              fr.get('pages', 0) >= 2, fr.get('pages'))
        shot(page3, 'c1-mobile-preview.png')
        check('мобайл: 0 JS-ошибок', len(errors3) == 0, errors3[:4])
        ctx3.close()

        browser.close()

    print('\nИТОГО: %d/%d (fail=%d)' % (PASS, PASS + FAIL, FAIL))
    return 1 if FAIL else 0


if __name__ == '__main__':
    import sys
    sys.exit(main())
